"""Build a source-backed DOCX and UTF-8 thesis-writing handoff from saved runs."""
from pathlib import Path
import csv, json, hashlib, shutil
from collections import Counter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
PROJECT = ROOT / 'calibration_gui'
RUNS = PROJECT / 'analysis_outputs/live_sensor_manual_only/model_runs'
B = RUNS / 'retraining_20260909'
T = RUNS / 'tuning_20260909'
def load(p): return json.loads(p.read_text(encoding='utf-8'))
def csvrows(p):
    with p.open(encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))
base, tuned = csvrows(B/'model_comparison.csv'), csvrows(T/'model_comparison.csv')
def row(rows,task,model): return next(r for r in rows if r['task']==task and r['model']==model)
def val(r,k): return float(r[k])
def n(x): return f'{float(x):.3f}'
def pct(x): return f'{float(x)*100:.2f}%'
WIN = {'force':'Extra Trees','localization':'Logistic regression','contact':'Random Forest'}
protocol=load(T/'protocol.json'); configs={c['id']:c for c in protocol['candidates']}
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(10.5 if name=='Normal' else {'Title':25,'Subtitle':12,'Heading 1':17,'Heading 2':12,'Heading 3':11,'Caption':9}[name])
    s.paragraph_format.space_after=Pt(6)
    if name.startswith('Heading'):s.paragraph_format.space_before=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.core_properties.title='Manual Calibration Models Technical Report for Thesis Writing'
doc.core_properties.subject='Force estimation contact detection and nine ROI localization'
doc.core_properties.author=''
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Manual calibration models  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for r in foot.runs:r.font.size=Pt(8)
text=[];tables=0
def p(s,style=None):doc.add_paragraph(s,style);text.append(s+'\n')
def h(s,level=1):doc.add_heading(s,level);text.append('\n'+s+'\n'+'='*len(s)+'\n')
def page(s):doc.add_page_break();h(s)
def table(title,headers,rows,widths=None):
    global tables
    tables+=1;p(f'Table {tables}. {title}','Caption')
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [6.77/len(headers)]*len(headers)
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,s in enumerate(headers):t.rows[0].cells[i].text=s
    for values in rows:
        cells=t.add_row().cells
        for i,s in enumerate(values):cells[i].text=str(s)
    for ri,r in enumerate(t.rows):
        trpr=r._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(r.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'233747' if ri==0 else ('F1F4F6' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.03
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(9);run.font.bold=(ri==0);run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    text.append('\t'.join(headers)+'\n'+'\n'.join('\t'.join(map(str,r)) for r in rows)+'\n')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

p('Manual Calibration Models Technical Report for Thesis Writing','Title')
p('Force estimation contact detection and nine ROI localization','Subtitle')
p('Prepared 5 October 2026  |  Experiments recorded 9 September 2026')
h('Executive summary')
p('Manual calibration data were used to retrain and compare force regression, binary contact detection, and nine-region localization models. A fixed-configuration benchmark was followed by a nested, grouped search over 178 model and feature configurations. The search also evaluated causal temporal smoothing, force residual corrections, contact decision thresholds, and ensembles. The strongest defensible headline is the performance of the complete nested selection procedure on the same held-out recordings used in the benchmark. [S1–S5]')
summary=[]
for task,metric,label,formatter in [('force','mae_N','Force MAE',n),('force','rmse_N','Force RMSE',n),('force','r2','Force R squared',n),('localization','accuracy','Localization accuracy',pct),('localization','macro_f1','Localization macro F1',pct),('contact','balanced_accuracy','Contact balanced accuracy',pct),('contact','recall','Contact recall',pct),('contact','fpr','Contact false positive rate',pct)]:
    a=row(base,task,WIN[task]);b=row(tuned,task,'Tuned selected');summary.append([label+(' (N)' if metric.endswith('_N') else ''),formatter(a[metric]),formatter(b[metric])])
table('Overall held-out results with equal recording weights',['Metric','Previous best fixed model','Tuned nested selector'],summary,[2.9,1.93,1.94])
p('The force MAE reduction is 21.68%. Localization accuracy increases by 10.62 percentage points and contact balanced accuracy by 6.64 percentage points. These are retrospective improvements within the archived manual dataset. The pooled improvement does not extend to every subgroup: force error increases in all four bins below 2 N, ROI 2 localization recall declines, and TEST1 remains substantially weaker than the other groups.')
p('The trained bundles are research artifacts. They were saved for offline and stateful inference, but were not substituted into the live GUI. No prospective validation, continuous localization error in millimetres, validated operating range, or deployed response-delay result is established by this study.')

page('1 Purpose and scope')
p('This report supplies the technical evidence needed to write thesis methodology, results, discussion, and conclusions for the manual-only model experiments. It documents what was implemented, what was measured, and which conclusions the measurements support. The report separates the September retraining and tuning experiments from earlier hardware validation and application development.')
h('Research objectives',2)
p('The objectives were to estimate applied force in newtons from optical measurements, identify whether contact is present, classify the contacted region among nine fixed ROIs, compare alternative model families under a consistent validation protocol, and improve predictive performance without introducing automatic-calibration data or target leakage.')
h('Task definitions',2)
table('Prediction targets and principal evaluation criteria',['Task','Target','Primary selection metric'],[['Force estimation','Continuous reference force in N','Minimum session-balanced MAE'],['Contact detection','0 dedicated no contact; 1 eligible press','Maximum balanced accuracy'],['Localization','ROI class 1 through 9','Maximum macro F1']],[1.4,3.0,2.37])
p('Force estimation is regression; therefore accuracy is not its principal metric. MAE, RMSE, bias, and R squared quantify its error. Localization accuracy measures correct discrete regions, not distance from a true spatial coordinate. Contact detection evaluates individual eligible frames, not contact episodes or event onset times.')
h('Experimental chronology',2)
p('The project records physical camera and load-cell tests on 3 August 2026. The manual feature-store report is dated 13 August 2026. The fixed benchmark and expanded tuning runs were generated on 9 September 2026. This report assembles those saved results on 5 October 2026; it does not represent a new collection of calibration data or a new training run. [S1, S3, S6, S7]')
p('Earlier project reports contain other evaluation units and model variants, including event-level or selectively accepted predictions. Those values must not be combined with the present frame-level metrics as though they describe the same experiment. The primary comparison in this report fixes the held-out frame identities, target alignment, force interval, and recording weights.')

page('2 System and acquisition context')
p('The SPARE calibration application combines a camera stream, an Arduino Nano and HX711 load-cell stream, nine fixed optical ROIs, unloaded baseline capture, synchronization, and export of video and tabular data. The application also supports a separate printer-driven acquisition path, but automatic-calibration archives are excluded from the model study described here. The load cell supplies training and evaluation reference force; the optical model does not consume load-cell force at inference. [S6, S8]')
p('The project hardware record identifies the experimental camera as an Arducam IMX179 Camera Module. Its 3 August physical test used an Arduino Nano/HX711 serial connection at 115200 baud, with HX711 DOUT on D4 and SCK on D5. These are historical hardware validation observations. They do not establish that every archived manual session used identical camera controls or a traceable reference-force uncertainty budget. [S7]')
h('Authoritative image and orientation',2)
p('The immutable feature store used session_video.mp4 as the quantitative image source. Overlay, processed, and motion-magnified videos were not used for model features. Already oriented 480 by 640 original frames were not transformed a second time. The configured transformation for a 640 by 480 input is a 90 degree clockwise rotation followed by horizontal mirroring. The quantitative intensity channel is the OpenCV HSV value channel. [S3, S8]')
h('Optical response extraction',2)
p('The feature contract defines an unloaded per-pixel median value-channel baseline. At each pixel, the signed change is current V minus baseline median V; the positive change is the larger of this difference and zero. ROI summaries retain positive light, signed light, active-pixel fractions, pixel statistics, and local intensity centroids. Signed features preserve decreases in intensity that a positive-only response would discard. The model experiments consume the saved feature partitions rather than reprocessing annotated or enhanced videos. [S3, S8, S9]')
p('The richer model features use per-ROI area, signed median and median absolute deviation, raw intensity mean/median/maximum, local centroid coordinates, and ROI geometry. A centroid missing from the optical calculation is represented by a centered coordinate together with a presence flag. The geometry normalizes optical centroids within each ROI; it does not supply the true contacted ROI to the estimator.')
h('Evidence boundary',2)
p('A camera acquisition throughput result from a separate hardware test cannot be substituted for the cadence of archived optical features or for end-to-end model latency. Likewise, the ability to record a physical session does not by itself validate the resulting force model. Hardware operation, offline prediction, and deployed sensor acceptance are distinct evidence categories.')

page('3 Dataset provenance and eligibility')
p('Only Manual Calibration (2).zip was permitted. Its recorded SHA256 is e039dcd49f3390d34dc979490dcf46a6ecc43d9e6a2aec08e4952acc14b5006a. The source firewall rejects the automatic-calibration archive and its derived lineage. Training read the immutable manual feature store feature-store-fa2500440adab6ca; partition and provenance checks tie those features to the permitted archive. [S3, S8, S10]')
table('Manual archive composition',['Recording role','Recordings','Frames','Use'],[['Primary press','54','24,453','Eligible frames for all tasks'],['Dedicated no contact','11','1,352','Contact target 0 and train-only floors'],['Replay only','18','3,513','Excluded from fitting and scoring'],['Total','83','29,318','Manual archive only']],[1.65,.8,.85,3.47])
p('Each TEST group contains one primary recording for each of the nine ROIs. Six TEST groups therefore contain 54 primary press recordings. Complete dedicated no-contact recordings are assigned to groups, with two in each of TEST1 through TEST5 and one in TEST6. Assignment keeps each recording intact across training and validation.')
coverage=csvrows(B/'evaluation_coverage.csv')
table('Primary-frame eligibility after fold-specific alignment',['Category','Frames'],[['All primary frames',sum(int(r['primary_frames']) for r in coverage)],['No valid aligned force',sum(int(r['unpaired_frames']) for r in coverage)],['Aligned force below 0.05 N',sum(int(r['below_005_N']) for r in coverage)],['Aligned force above 3 N',sum(int(r['above_3_N']) for r in coverage)],['Evaluated press frames',sum(int(r['evaluated_frames']) for r in coverage)]],[5,1.77])
p('Force and localization use 7,154 eligible press frames from 54 recordings. Contact detection adds all 1,352 dedicated no-contact frames, giving 8,506 frames from 65 recordings. Press frames outside the exploratory 0.05–3 N interval are excluded rather than relabelled as no contact. Thus the binary task distinguishes eligible press frames from dedicated no-contact recordings; it does not validate contact status for every excluded frame.')
p('The interval is an exploratory analysis rule. Its lower endpoint is not an established detection limit, and its upper endpoint is not a validated maximum load. Most primary frames lie above the interval; the results therefore describe a restricted part of the archived acquisition, not the whole force distribution.')

page('4 Alignment preprocessing and leakage controls')
h('Timestamp alignment',2)
p('For optical time t and lag L, the target is reference force at t minus L. Force is linearly interpolated within the same recording. Duplicate reference timestamps are reduced by the median, extrapolation is prohibited, and interpolation across gaps greater than 300 ms is rejected. A positive lag therefore associates the current optical response with an earlier reference-force measurement. This convention must be preserved when interpreting delay. [S11]')
p('The benchmark selected lag from 0, 100, 200, 300, 400, and 500 ms using inner grouped NNLS MAE, taking the smallest lag within 0.5% of the best error. The outer-fold lags for TEST1–TEST6 were 300, 300, 200, 300, 300, and 300 ms. The tuning experiment reused the same training-subset lag decisions so the compared outer rows and targets remained identical. The broader application configuration contains other settings; the saved run protocol and executed training source govern this experiment. [S1, S2, S9]')
h('Training-only floor and scale estimation',2)
p('For each ROI, the optical floor is the median of the within-recording medians of raw positive-light sums from training no-contact recordings. The corrected positive feature is max(raw positive sum minus fitted floor, 0) divided by ROI pixel area. Signed-light sums are divided by area without positive clipping. A weighted StandardScaler is fitted only to eligible training rows. Centering is disabled for positive-only force features to preserve nonnegativity and enabled for the other packs.')
p('Floor and scaler parameters are refitted at each training boundary. No held-out no-contact recording contributes to the floor. The immutable feature store contains optical measurements but no globally fitted floor, response scale, decision threshold, imputation, detection limit, maximum force, or model transform. This separation prevents test information from entering trainable preprocessing.')
h('Prohibited inference inputs',2)
p('Reference force, true ROI label, TEST number, recording identity, absolute time, and recording duration are not estimator features. Recording identifiers and relative timestamps are used for partitioning and resetting causal history only. The inference audit confirms that changing or omitting the target columns does not change predictions. Future-frame perturbations do not alter earlier outputs. [S10]')
p('Causal inference may use earlier optical frames from the same held-out recording, including earlier frames outside the evaluated force interval. Their labels are not used. Consequently, these scores describe chronological sequence inference with available past optical history; they should not be represented as independent cold-start predictions on every selected frame.')

page('5 Validation design and metric definitions')
p('The outer loop holds out one complete TEST group and its assigned no-contact recordings. The remaining five groups supply model development. Within each outer training set, five inner folds fit four groups and validate the remaining group. Hyperparameters, feature pack, smoothing, residual correction, family, ensemble, and contact threshold are chosen from these inner predictions before the outer group is scored. This produces six disjoint outer test groups while preserving correlation within each recording. [S1, S2, S9]')
p('The fifteen possible four-group training subsets are cached once. A cache may store predictions for both excluded groups, but selection for a given outer fold reads only the appropriate inner-validation group. The outer test group is excluded from fitting and selection. The audit verifies 30 inner train/validation boundaries. After outer evaluation, the final research models are selected using all fifteen leave-two-groups-out splits and refitted to all eligible manual training data. Their all-data fit is not a held-out performance result.')
h('Equal recording weights',2)
p('If recording s has n eligible rows in the metric being computed, each row receives weight proportional to 1/n. Thus every contributing recording has equal total weight. The implementation rescales weights to mean one, which does not alter weighted averages. A bin-specific calculation recomputes weights among rows in that bin, so recordings with no rows in a bin do not contribute. These are not unweighted frame averages.')
h('Regression metrics',2)
p('With error e equal to predicted force minus reference force, MAE is the weighted mean of absolute e, RMSE is the square root of the weighted mean of squared e, and bias is the weighted mean of e. Positive bias indicates overestimation. R squared is one minus weighted squared error divided by the weighted squared deviation of reference force from its weighted evaluation mean. Negative R squared means worse squared error than that evaluation-mean reference; the mean is not a separately trained deployment model.')
h('Classification metrics',2)
p('The weighted confusion matrix sums recording-balanced weights for each true and predicted class. Accuracy is its diagonal sum divided by total weight. Recall is TP/(TP+FN), precision is TP/(TP+FP), and false-positive rate is FP/(FP+TN). Binary balanced accuracy is the mean of contact recall and no-contact specificity. Class F1 is 2TP/(2TP+FP+FN); macro F1 averages class F1 across all nine ROI classes or both contact classes. Undefined divisions are handled consistently as zero in the saved calculations.')
p('Pooled metrics are computed after joining all outer predictions; pooled macro F1 need not equal the average of six fold-level macro F1 values. Contact precision also depends on the archive’s contact/no-contact mixture. No pooled tuned AUC is reported because adaptive families and uncalibrated score scales differ between folds.')

page('6 Fixed model benchmark')
p('The initial benchmark used seed 1729 and fixed configurations specified before the run’s outer predictions were inspected. Nine force configurations included simple reference models, constrained regressors, nonlinear tree models, and support-vector regression. Six classifier families were evaluated separately for contact and localization. Twenty-two distinct training subsets, including inner subsets, outer training sets, and the full pool, produced 462 fitted estimators. [S1, S9]')
p('Force models used nine floor-corrected positive-light features divided by ROI area. Total-light and isotonic baselines sum the scaled channels. Classifiers used 27 channels comprising positive light, signed light per pixel, and active fractions. Force fitting balanced recordings and 0.25 N force bins; localization balanced recordings; contact fitting balanced classes and recordings. Evaluation always used equal recording weights. No temporal smoothing or threshold tuning was applied.')
table('Fixed force configurations and held-out results',['Model','MAE N','RMSE N','Bias N','R squared'],[[r['model'],n(r['mae_N']),n(r['rmse_N']),n(r['bias_N']),n(r['r2'])] for r in base if r['task']=='force' and r['model']!='Nested selected'],[2.45,1.08,1.08,1.08,1.08])
p('Extra Trees was the best fixed force model and was selected in every outer inner search. Its 0.637 N MAE improved on the fitted median baseline, but negative R squared and negative bias showed limited tracking and systematic underestimation. The monotonic models did not outperform Extra Trees in this comparison. Poor linear performance does not prove that all physically constrained models fail; it applies to the tested features, interval, configurations, and data.')
table('Fixed classifier comparison',['Family','ROI accuracy','ROI macro F1','Contact balanced accuracy','Contact FPR'],[[m,pct(row(base,'localization',m)['accuracy']),pct(row(base,'localization',m)['macro_f1']),pct(row(base,'contact',m)['balanced_accuracy']),pct(row(base,'contact',m)['fpr'])] for m in ['Logistic regression','RBF SVC','Random Forest','Extra Trees','HistGB','XGBoost']],[1.85,1.17,1.17,1.4,1.18])
p('Logistic regression narrowly led localization by the prespecified macro-F1 criterion, while RBF SVC had slightly higher accuracy. Random Forest led contact balanced accuracy but produced an 8.20% no-contact false-positive rate. These fixed-family winners were identified by comparing outer results; they must be distinguished from a family selector evaluated wholly inside training.')

page('7 Expanded model and feature search')
p('The tuning run used seed 1741 and a frozen finite grid of 178 configurations: 62 force, 64 localization, and 52 contact configurations. Fifteen four-group subsets produced 2,670 inner model fits. Additional choices of output smoothing, residual offset, threshold, and ensemble were evaluated using cached inner predictions. The count of 178 refers to estimator and feature configurations, not every derived policy combination. [S2, S9]')
table('Optical feature packs',['Code name','Dimensions','Contents'],[['positive','9','Floor-corrected positive light per pixel'],['optical','27','Positive light, signed light, active fractions'],['shape','90','Optical plus normalized and centered spatial patterns'],['spatial','180','Shape plus pixel statistics, centroids and intensity transforms'],['invariant','117','Normalized patterns and local spatial statistics'],['temporal','288','Spatial plus causal histories and differences']],[1.1,.85,4.82])
p('Across ROIs, L1 normalization divides a vector by its sum of absolute values plus 0.000001. Median centering subtracts the across-ROI median. Shape features include normalized positive light, normalized centered signed light, normalized active fractions, centered signed/active patterns, normalized raw-mean contrast, and normalized signed-median contrast. The spatial pack adds signed median, median absolute deviation, normalized centroid coordinates, centroid presence, raw intensity contrasts, maximum-to-mean ratios, and logarithmic transforms.')
p('The code name invariant denotes a lighting-normalized candidate pack. It does not establish mathematical or empirical invariance to every illumination, exposure, geometry, or sensor change. Several of its channels retain magnitude information. Any thesis claim of illumination invariance would require a separate controlled evaluation.')
h('Causal temporal processing',2)
p('Temporal input features append exponential moving averages of the 27 optical channels with 0.3 s and 1.0 s time constants, the current-minus-fast difference, and the fast-minus-slow difference. At interval dt, alpha equals 1 minus exp(−dt/tau), and the new state equals the previous state plus alpha times the current-minus-previous difference. State is initialized from the current row at a new recording, a gap greater than 0.3 s, or a nonpositive interval.')
p('Output smoothing separately tests time constants 0, 0.3, and 1.0 s. Zero means no smoothing. Causality prevents future information leakage but does not eliminate response delay. Because the true ROI stays fixed within each press recording, temporal gains cannot establish reliable tracking while contact moves rapidly between ROIs.')

page('8 Hyperparameters corrections and selection policies')
table('Estimator search ranges',['Task and family','Varied settings'],[['Force Extra Trees','Leaf size and feature fraction: (3, 0.75), (15, 0.75), (30, 1.0); positive-only baseline variants also included'],['Force Random Forest','Minimum leaf size 5 or 20; feature fraction 0.75'],['Force HistGB','15 or 31 leaves; squared loss, plus 15-leaf absolute loss; L2 10'],['Force XGBoost','Depth 3 or 5; lambda 20'],['Force RBF SVR','C 2 or 20; epsilon 0.05; gamma scale'],['Force MLP','Optical or spatial pack; hidden layers 64 and 32; alpha 1'],['ROI logistic regression','C 0.03, 0.3, 3, or 30 across five packs'],['ROI RBF SVC','C 1, 10, or 100; gamma multiplier 0.1 or 1'],['ROI Extra Trees','Leaf size 2 or 10; feature fraction 0.75'],['ROI HistGB and MLP','Spatial or invariant pack; 15 leaves and L2 10, or 64 and 32 hidden units'],['Contact tree models','RF leaf 2, 10, 30; ET leaf 3, 15; HistGB leaves 7, 15; XGBoost depth 2, 4'],['Contact linear and kernel models','Logistic C 0.1 or 10; RBF SVC C 1 or 10 with gamma multiplier 1']],[1.7,5.07])
p('Common settings were 180 estimators for forests and Extra Trees; 180 iterations, learning rate 0.05, minimum leaf size 20, and disabled early stopping for HistGB; and 220 estimators, learning rate 0.04, minimum child weight 5, histogram tree method, and one estimator thread for XGBoost. Logistic regression allowed 2,500 iterations. MLP used 350 iterations, learning rate 0.001, and no early stopping. RBF SVC gamma is the selected multiplier divided by the number of features. The full exact grid is retained in protocol.json and in the evidence package.')
p('Force predictions are nonnegative. Each inner candidate compares no offset, a weighted mean residual offset, and a weighted median residual offset, followed by nonnegative clipping. Most force candidates use equal-recording training weights; the positive-only Extra Trees bin-weighted candidate preserves a comparison with the earlier weighting rule.')
p('For localization, the winning class maximizes the predicted class score. SVC decision outputs are mapped with softmax for localization and a sigmoid for contact. These transformations are score mappings, not demonstrations of calibrated probabilities. For contact, every exact weighted ROC operating point is available for training-only threshold selection. The main policy maximizes balanced accuracy; the alternative maximizes recall subject to inner FPR at most 5%.')
p('For force and localization, the best candidate from each family is identified first. The three best distinct families are averaged with equal weights, and this ensemble replaces the single-family selection only if it improves the inner primary metric. Contact uses a selected single candidate and threshold. An ensemble is not automatically best in every outer fold.')

page('9 Overall results and fair comparisons')
p('All rows below use the same outer frame identities and truth values. Previous best fixed model means the family ranked best after the initial outer comparison. Previous nested selector means the original procedure chose its family using inner validation. Tuned nested selector means the expanded search chose its complete policy inside outer training. Current frame only is a separate inner-selected pipeline excluding temporal input features and output smoothing. [S4, S5]')
for task,columns in [('force',[('mae_N','MAE N'),('rmse_N','RMSE N'),('bias_N','Bias N'),('r2','R squared')]),('localization',[('accuracy','Accuracy'),('macro_f1','Macro F1')]),('contact',[('accuracy','Accuracy'),('balanced_accuracy','Balanced accuracy'),('recall','Recall'),('fpr','FPR'),('macro_f1','Macro F1')])]:
    pairs=[('Previous best fixed',row(base,task,WIN[task])),('Previous nested',row(base,task,'Nested selected')),('Tuned nested',row(tuned,task,'Tuned selected')),('Current frame only',row(tuned,task,'Tuned frame-only'))]
    table(task.capitalize()+' held-out comparison',['Procedure']+[c[1] for c in columns],[[label]+[(n if task=='force' else pct)(r[c[0]]) for c in columns] for label,r in pairs],[1.9]+[(6.77-1.9)/len(columns)]*len(columns))
p('Relative to the previous best fixed models, force MAE fell from 0.637 to 0.499 N and RMSE from 0.770 to 0.632 N. Bias moved from −0.289 to +0.078 N. R squared increased from −0.030 to 0.306, indicating improved but still incomplete explanation of force variation. These values do not support describing the force estimator as highly accurate across the entire interval.')
p('Localization accuracy rose from 73.79% to 84.40%, and macro F1 from 74.00% to 84.75%. Contact precision rose from 98.22% to 100.00% on the recorded evaluation mixture. The observed zero FPR is based on 1,352 no-contact frames from only 11 recordings, with temporal dependence within each recording. It is not proof of zero false-alarm risk in deployment.')
p('The current-frame pipeline also improved the principal metrics relative to the earlier fixed winners. This shows that the measured overall improvement is not confined to a pipeline requiring history. It does not isolate the contribution of smoothing: feature choice, model family, hyperparameters, and ensemble policy also differ between the selected pipelines.')

page('10 Tuned model family comparisons')
p('Each family row below tunes that family within each outer training set. The selected-family row may change family across folds. Ranking these rows after observing outer results is a retrospective comparison, so the best family row must not replace the nested-selector headline without acknowledging the additional selection. [S4]')
table('Tuned force families',['Family or policy','MAE N','RMSE N','R squared'],[[r['model'].replace('Tuned ',''),n(r['mae_N']),n(r['rmse_N']),n(r['r2'])] for r in sorted([x for x in tuned if x['task']=='force' and x['model']!='Tuned frame-only'],key=lambda x:float(x['mae_N']))],[2.8,1.32,1.32,1.33])
table('Tuned localization families',['Family or policy','Accuracy','Macro F1'],[[r['model'].replace('Tuned ',''),pct(r['accuracy']),pct(r['macro_f1'])] for r in sorted([x for x in tuned if x['task']=='localization' and x['model']!='Tuned frame-only'],key=lambda x:-float(x['macro_f1']))],[3.17,1.8,1.8])
p('The force ensemble has the lowest retrospective family-policy MAE, 0.497 N, only slightly below the nested selector’s 0.499 N. For localization, the ensemble leads macro F1 at 85.12%, while tuned logistic regression has 84.88% accuracy and 85.00% macro F1. The evidence supports these as strong research candidates but does not establish a statistically decisive separation among the closely ranked alternatives.')

page('11 Contact operating point comparisons')
table('Tuned contact families and policies',['Family or policy','Balanced accuracy','Recall','FPR','Precision'],[[r['model'].replace('Tuned ',''),pct(r['balanced_accuracy']),pct(r['recall']),pct(r['fpr']),pct(r['precision'])] for r in sorted([x for x in tuned if x['task']=='contact'],key=lambda x:-float(x['balanced_accuracy']))],[2.15,1.22,1.13,1.13,1.14])
p('The retrospectively strongest contact family is tuned RBF SVC, with 99.54% balanced accuracy, 99.09% recall, and zero observed FPR. The main nested selection procedure obtains 98.50% balanced accuracy, 97.00% recall, and zero observed FPR. The difference reflects the distinction between always using a particular tuned family and allowing training data to select among families. The full-data training search ultimately chose an RBF SVC for the saved primary contact model.')
h('Why the FPR constrained alternative is not the default',2)
p('The alternative policy enforces FPR at most 5% on inner validation predictions only. On outer recordings, it reaches 99.57% recall but 8.76% FPR and 95.40% balanced accuracy. Its higher overall accuracy and macro F1 reflect a different tradeoff; they do not mean it satisfies the intended false-alarm constraint on new recordings. The main balanced-accuracy policy is the selected research default.')
p('Thresholds are selected using training-only out-of-fold scores. They are not universal constants transferable to a different feature extractor, model fit, score transformation, or sampling regime. The final saved threshold applies to its associated saved estimator and preprocessing. Threshold tuning must be repeated within training if those components change.')
h('Interpretation for sensor operation',2)
p('Frame-level false positives and recall do not characterize the duration or rate of false-contact episodes. A small number of correlated false frames can produce one episode or several episodes depending on debounce. Similarly, smoothing can make a contact easier to classify during a sustained press while delaying onset or release. Those effects require event-level measurements in a prospective live experiment.')

page('12 Variation across TEST groups and ROIs')
bfold=csvrows(B/'fold_metrics.csv');tfold=csvrows(T/'fold_metrics.csv')
def fold(rows,task,model,i):return next(r for r in rows if r['task']==task and r['model']==model and int(r['outer_fold'])==i)
table('Localization accuracy by complete held-out group',['Group','Previous logistic','Tuned selector','Current frame'],[[f'TEST{i}',pct(fold(bfold,'localization','Logistic regression',i)['accuracy']),pct(fold(tfold,'localization','Tuned selected',i)['accuracy']),pct(fold(tfold,'localization','Tuned frame-only',i)['accuracy'])] for i in range(1,7)],[1.1,1.89,1.89,1.89])
p('TEST1 remains the largest localization weakness. Its tuned causal accuracy is 42.71%, compared with 87.43–97.87% for the other five groups. The current-frame pipeline performs better on TEST1 but has lower pooled performance. Because these pipelines differ in several selected components, the result cannot be assigned to smoothing alone. The original report observed different raw brightness in TEST1; that association does not establish the cause of the generalization gap.')
br=csvrows(B/'roi_metrics.csv');tr=csvrows(T/'roi_metrics.csv')
def roi(rows,task,i):return next(r for r in rows if r['task']==task and int(r['roi'])==i and ('model' not in r or r['model']==WIN[task]))
table('Force and localization results by ROI',['ROI','Frames','Baseline MAE N','Tuned MAE N','Baseline ROI recall','Tuned ROI recall'],[[i,roi(tr,'force',i)['frames'],n(roi(br,'force',i)['mae_N']),n(roi(tr,'force',i)['mae_N']),pct(roi(br,'localization',i)['recall']),pct(roi(tr,'localization',i)['recall'])] for i in range(1,10)],[.48,.68,1.27,1.27,1.53,1.54])
p('Every ROI is represented by six press recordings. ROI 2 has tuned localization recall of 56.00%, lower than its baseline recall, despite the overall gain. ROI 9 is strongest at 99.67%. The weakest tuned force MAE is 0.606 N in ROI 1 and the lowest is 0.410 N in ROI 9. Per-ROI scores are conditional on the evaluated force interval and do not establish a continuous spatial resolution.')

page('13 Force range diagnostics and combined predictions')
bb=csvrows(B/'force_bin_metrics.csv');tb=csvrows(T/'force_bin_metrics.csv')
table('Force error by diagnostic reference-force bin',['Force bin','Frames','Recordings','Baseline MAE N','Tuned MAE N'],[[r['force_bin'],r['frames'],r['sessions'],n(a['force_mae_N']),n(r['force_mae_N'])] for a,r in zip(bb,tb)],[1.45,.72,.85,1.87,1.88])
p('Force MAE worsens in all four diagnostic bins below 2 N, including 1.107 to 1.241 N below 0.5 N. It improves in the two upper bins, particularly 2.5–3 N, where MAE falls from 0.916 to 0.533 N. The lower four bins contain 3,527 frames and the upper two contain 3,627, but headline metrics use recording weights rather than these raw proportions. The pooled improvement must not be described as uniform improvement across force levels.')
table('Detection and localization by force bin',['Force bin','Baseline contact recall','Tuned contact recall','Baseline ROI accuracy','Tuned ROI accuracy'],[[r['force_bin'],pct(a['contact_recall']),pct(r['contact_recall']),pct(a['localization_accuracy']),pct(r['localization_accuracy'])] for a,r in zip(bb,tb)],[1.37,1.35,1.35,1.35,1.35])
p('Contact recall and localization accuracy improve in every listed bin, but below 0.5 N they remain only 85.41% and 57.31%, respectively. That bin contains just 263 frames spread over 52 recordings. The selected objective, changed training weights, richer features, and temporal policies may influence these tradeoffs, but this experiment does not provide an isolated causal ablation for each factor.')
be,te=load(B/'end_to_end_metrics.json'),load(T/'end_to_end_metrics.json')
table('Combined pipeline on eligible positive frames',['Metric','Previous combination','Tuned combination'],[[label,fmt(be[k]),fmt(te[k])] for k,label,fmt in [('mae_N','Contact-gated force MAE N',n),('rmse_N','Contact-gated force RMSE N',n),('joint_contact_and_roi_accuracy','Correct contact and ROI',pct)]],[3.1,1.83,1.84])
p('For the combined force metric, a missed contact is assigned zero force. Joint accuracy requires both a contact decision and the correct ROI. These combined metrics use the 7,154 eligible positive frames and equal recording weights. They do not quantify force outputs on dedicated no-contact recordings or performance outside the evaluation interval.')

page('14 Uncertainty and interpretation')
unc=load(T/'paired_uncertainty.json')
table('Descriptive paired changes across six outer groups',['Metric','Mean fold change','Bootstrap 95 percent interval'],[[r['task']+' '+r['metric'],f"{r['mean_fold_change']*(1 if r['task']=='force' else 100):+.3f}"+(' N' if r['task']=='force' else ' pp'),f"{r['bootstrap95_low']*(1 if r['task']=='force' else 100):+.3f} to {r['bootstrap95_high']*(1 if r['task']=='force' else 100):+.3f}"+(' N' if r['task']=='force' else ' pp')] for r in unc],[2.15,1.67,2.95])
p('Changes are tuned selector minus previous best fixed family. A negative MAE change and positive macro-F1 or balanced-accuracy changes indicate improvement. All six paired group changes have the favorable sign for these three metrics. The intervals summarize resampling of six observed group pairs; they are descriptive rather than proof of prospective generalization. Cross-validation training sets overlap, the number of groups is small, and the recordings were already used in prior development. [S12]')
p('The mean fold change for localization macro F1 is 9.71 percentage points, whereas the pooled macro-F1 change is 10.76 percentage points. The difference is expected because F1 is nonlinear and its aggregation differs. Contact mean fold balanced-accuracy change also differs from the pooled change because the distribution of no-contact recordings is unequal across the six groups. A thesis should retain the aggregation label beside each estimate.')
h('What improved and what remains unresolved',2)
p('The expanded search improves pooled regression error, discrete localization, and contact discrimination on the controlled comparison. Current-frame improvements suggest that feature and model selection can provide useful gains even when history is unavailable. Force estimation remains the weakest quantitative component: R squared is 0.306 and the large low-force errors limit claims about light contacts. The contact task is strong on archived dedicated no-contact sequences, but its short and narrow negative-data coverage limits generalization.')
p('The analysis supports a statement that the complete tuning procedure improved overall retrospective performance. It does not identify a single technique as the cause of all improvement. To isolate feature normalization, temporal input, output smoothing, weighting, or ensembling, a new experiment would need matched ablations with those factors changed one at a time inside the same validation boundaries.')
h('Meaning of best model',2)
p('For this dataset, a force ensemble and a localization ensemble are the selected final research artifacts, while contact uses RBF SVC. These choices reflect training-only selection on the complete pool after the nested evaluation. Best means preferred under the stated objectives and candidate search; it does not mean universally best, optimal for live latency, or validated for unseen operating conditions.')

page('15 Saved models and inference behavior')
finalrows=[]
for task in ['force','localization','contact']:
    spec=load(T/'models'/f'{task}.json')
    for member in spec['members']:
        choice=member['choice'];c=configs[choice['id']]
        finalrows.append([task,c['family'],c['features'],str(choice['tau']),json.dumps(c['params'],sort_keys=True)])
table('Final full-pool model components',['Task','Family','Feature pack','Output tau s','Varied parameters'],finalrows,[.85,1.18,.82,.73,3.19])
p('Each ensemble gives its three members equal weight after their member-specific corrections or smoothing. The force members are HistGB, Extra Trees, and XGBoost using temporal inputs. Localization uses HistGB, Extra Trees, and logistic regression with the invariant pack and 1 s output smoothing. Contact uses the shape pack, RBF SVC with C 1 and gamma 1/90, and 1 s output smoothing. [S13]')
details=[]
for task in ['force','localization','contact','contact_fpr5']:
    path=T/'models'/f'{task}.json'
    if not path.exists():continue
    spec=load(path)
    for m in spec['members']:
        c=m['choice'];tail=[]
        for k in ['correction','offset','threshold','threshold_policy']:
            if k in c:tail.append(f'{k}={c[k]}')
        if tail:details.append([task,c['id'],'; '.join(tail)])
table('Final learned corrections and thresholds',['Task','Candidate','Saved value'],details,[1.1,.9,4.77])
p('The primary contact threshold is 0.32685983180999756 on its smoothed sigmoid-transformed SVC score. This number is a decision rule, not a calibrated probability of contact. Metrics stored in the final model choice JSON describe full-pool selection predictions and must not be reported as independent test accuracy.')
h('Inference contract',2)
p('The models directory contains causal bundles; models_frame_only contains separate bundles requiring only the current frame. The helper load_models loads trusted locally generated joblib artifacts. infer_sequence processes each complete recording chronologically. For live-style sequential use, instantiate StreamingResearchPredictor once, call predict_one for ordered optical rows, retain its state between frames, and reset it at a new recording. Gaps above 300 ms reset history. Supplying a fresh predictor for every frame would change the causal model’s intended behavior.')

page('16 Runtime software and verification')
timing=load(T/'inference_timing.json');env=load(T/'environment.json');audit=load(T/'validation_audit.json')
table('Offline inference timing',['Measurement','Result'],[['Streaming median',f"{timing['streaming_p50_ms']:.2f} ms"],['Streaming 95th percentile',f"{timing['streaming_p95_ms']:.2f} ms"],['96-frame batch median',f"{timing['median_batch_ms']:.2f} ms"],['Amortized batch time per frame',f"{timing['amortized_ms_per_frame']:.3f} ms"]],[4.7,2.07])
p('Timing includes the three primary tasks, the alternate contact policy, optical tabular preprocessing, and causal state. It excludes camera acquisition, video decoding, pixel feature extraction, display, and contact onset delay. The batch result is amortized throughput, not per-frame streaming latency. The older baseline timing used a different three-model scope, so it should not be treated as a controlled speed comparison. [S14]')
table('Recorded training environment',['Component','Version'],[['Python','3.12.14'],['Operating system',env['platform']]]+[[k,v] for k,v in env['packages'].items()],[3.2,3.57])
p('Twenty automated tests passed across the existing manual-data firewall and preprocessing tests and five new tuning-feature tests. An independent audit recalculated 114 metrics, with maximum discrepancy 1.249 × 10^-15. It confirmed identical outer frame IDs and truth values, complete-session boundaries, source hashes, saved model checksums, finite scalers, and reproduction of TEST1 predictions from refitted selected models. [S10]')
p('Inference checks covered target/reference-column invariance, omission of labels, causal-prefix and future-perturbation invariance, session reset behavior, current-frame batch versus single-row parity, and stateful streaming versus full-sequence parity. These checks establish implementation consistency; they do not substitute for new physical validation data. The report delivery records also preserve source snapshots and model checksums for traceability.')

page('17 Limitations and next experimental steps')
h('Retrospective data reuse',2)
p('All six TEST groups had already been examined in earlier project development. The nested search controls fitting and selection leakage within this run, but it cannot restore an untouched prospective test set. Model improvements should therefore be described as retrospective manual-only evidence. A final claim about generalization requires new recordings collected after freezing the candidate model and evaluation rules.')
h('Coverage and task limitations',2)
p('The study contains six primary acquisition groups and 11 dedicated no-contact recordings. The evaluated force interval is restricted, low-force frames are sparse, and ROI is stationary within each press recording. Continuous location ground truth, dynamic transitions between ROIs, multi-contact force separation, sustained environmental variation, and long no-contact runs are not established by these metrics. The original operating-range and application acceptance gates have not been revalidated.')
h('Force model acceptance',2)
p('The selected nonlinear force models are not constrained to be monotonic and have not passed a new monotonicity or optical-support acceptance audit. Clipping predictions to nonnegative values does not establish physical validity. Positive pooled R squared is an improvement over the original benchmark, but large low-force errors and bin-specific degradation remain material limitations.')
h('Recommended prospective manual study',2)
p('Freeze the saved model, input contract, lag interpretation, decision thresholds, and reset rules before collecting a new manual-only test set. Record acquisition settings and baseline conditions. Balance coverage across ROIs and force levels, especially below 1 N; include loading and unloading, long no-contact intervals, altered lighting, and movement between regions. Keep all recordings from a session or acquisition block together when reserving the test set.')
p('Measure force MAE, RMSE, signed bias, and error quantiles by force bin and ROI; contact recall, false-contact episodes per unit time, release and onset delay; discrete ROI accuracy during transitions; and total acquisition-to-display latency. Document the reference load cell, calibration procedure, fixture, uncertainty, sensor geometry, and material properties in lab records. Evaluate a predefined operating range rather than choosing the range after looking at favorable errors.')
p('If low-force accuracy is a primary requirement, consider a new training objective that explicitly weights low-force error or a separately validated force-regime model. Such changes remain proposals, not completed results. They must be tuned inside training and evaluated on the frozen prospective split. New automatic-calibration data should remain outside this manual-only model study unless the scientific scope is explicitly changed.')

page('18 Guidance for thesis writing')
p('Use this report as technical source material for the model-development portion of the thesis. The following chapter mapping is a writing aid, not a replacement for the university’s prescribed format.')
table('Suggested thesis mapping',['Thesis part','Material to develop'],[['Introduction','Problem, three prediction tasks, objectives, and manual-only scope'],['Related literature','Independently sourced optical tactile sensing, force estimation, grouped validation, and model-selection literature'],['Methodology','Sections 2–8 and 15–16: apparatus context, dataset, preprocessing, validation, search and implementation'],['Results','Sections 9–14: pooled metrics, family comparisons, subgroup diagnostics, combined predictions and uncertainty'],['Discussion','Tradeoffs, data reuse, low-force errors, TEST1, ROI 2, threshold and temporal limitations'],['Conclusion and recommendations','Supported overall gains, offline research status, and prospective manual validation plan']],[1.7,5.07])
h('Supported statements',2)
p('The manual-only tuning procedure reduced held-out session-balanced force MAE from 0.637 to 0.499 N, increased nine-ROI classification accuracy from 73.79% to 84.40%, and increased contact balanced accuracy from 91.86% to 98.50% on the same retrospective evaluation frames. All three principal metrics improved across the six paired groups, while force-bin and ROI diagnostics revealed remaining weaknesses.')
h('Statements that would overstate the evidence',2)
p('Do not claim that every metric in every subgroup improved; the force bins below 2 N and ROI 2 contradict that claim. Do not call 0.05 N a validated detection limit, 3 N a validated maximum force, or 100% observed precision a guarantee. Do not describe ROI classification as millimetre localization, the full-data refit as an independent test, or the saved research model as deployed and validated. Do not attribute all gains to one feature or tuning technique without an ablation.')
h('Information to obtain from the researcher',2)
p('Before completing the thesis apparatus and acquisition sections, verify sensor fabrication and dimensions, ROI physical spacing, camera distance and illumination, archived camera settings, operator protocol, press geometry, load-cell specification and traceable uncertainty, calibration mass documentation, and required university format. Keep these as explicit research questions rather than filling them with invented specifications. Historical hardware notes can support context but cannot certify conditions in every manual recording.')
p('A separate UTF-8 prompt accompanies this report for use in ChatGPT Writing. It instructs the writing model to preserve units, evaluation units, evidence status, and limitations, and to flag missing laboratory details and literature references for the researcher.')

page('19 Reproducibility and source register')
p('Paths in this section are relative to the SOFTWARE CALIBRATION workspace. For compact references, B denotes calibration_gui/analysis_outputs/live_sensor_manual_only/model_runs/retraining_20260909 and T denotes the sibling tuning_20260909 directory. The evidence subfolder accompanying this report contains copies of key tabular and JSON records plus a SHA256 manifest. It does not duplicate the large feature store or trained model binaries.')
sources=[
('S1','B/protocol.json; B/estimator_parameters.json; B/lag_selection.json','Fixed benchmark protocol and exact estimator settings'),
('S2','T/protocol.json; T/frame_only_protocol.json','Frozen search grid and current-frame comparison rules'),
('S3','B/provenance.json; T/provenance.json; B/evaluation_coverage.csv','Archive lineage, feature-store provenance and row exclusions'),
('S4','T/model_comparison.csv; T/before_after.csv','Pooled held-out tuning and comparison results'),
('S5','B/model_comparison.csv; B/fold_metrics.csv','Baseline family and nested selection results'),
('S6','calibration_gui/README.md','System and acquisition application overview'),
('S7','calibration_gui/PHYSICAL_HARDWARE_VALIDATION_20260803.md','Historical physical camera and load-cell test record'),
('S8','calibration_gui/config/live_sensor_manual_only.json','Manual-only firewall and optical input contract'),
('S9','calibration_gui/scripts/retrain_manual_models.py; tune_manual_models.py; evaluate_manual_frame_only.py','Executed training, features and selection implementation'),
('S10','T/validation_audit.json; T/metric_audit.csv; T/report_delivery.json; calibration_gui/tests/test_tuned_manual_features.py','Metric, boundary, provenance and inference checks'),
('S11','calibration_gui/core/timestamp_alignment.py; calibration_gui/scripts/fit_manual_only_preprocessing.py','Timestamp pairing, train-only floors, weights and lag logic'),
('S12','T/paired_uncertainty.json; B and T fold_metrics.csv, roi_metrics.csv, force_bin_metrics.csv, end_to_end_metrics.json','Group uncertainty and conditional diagnostics'),
('S13','T/models/*.json; T/models_frame_only/*.json; T/final_selections.json','Final saved members, corrections, thresholds and manifests'),
('S14','T/inference_timing.json; T/environment.json; calibration_gui/scripts/predict_tuned_manual.py','Inference contract, runtime scope and recorded environment')]
for sid,path,meaning in sources:p(f'[{sid}] {meaning}. {path}.')
p('These are primary project records, not external literature citations. The thesis literature review should use independently verified scholarly sources. Archive checksums identify data provenance; they do not measure laboratory uncertainty. Saved source snapshots in T/source_snapshot preserve the training-era implementation if the current working tree subsequently changes.')
commands=[
'.venv/Scripts/python.exe -u -c "from scripts.tune_manual_models import main; main()" --jobs 8 --output analysis_outputs/live_sensor_manual_only/model_runs/tuning_repeat',
'.venv/Scripts/python.exe -m scripts.evaluate_manual_frame_only --run analysis_outputs/live_sensor_manual_only/model_runs/tuning_repeat',
'.venv/Scripts/python.exe -m scripts.audit_tuned_manual --run analysis_outputs/live_sensor_manual_only/model_runs/tuning_repeat',
'.venv/Scripts/python.exe -m scripts.report_tuned_manual --run analysis_outputs/live_sensor_manual_only/model_runs/tuning_repeat']
page('Appendix A Exact saved model selection records')
p('The following compact records identify the selected final members and current-frame alternatives. Candidate IDs refer to the frozen 178-entry protocol copied into the evidence folder. Selection scores embedded in the original JSON are intentionally omitted here to avoid confusing full-pool selection with independent outer evaluation.')
for folder,label in [('models','Causal models'),('models_frame_only','Current frame models')]:
    h(label,2)
    for task in ['force','localization','contact']:
        spec=load(T/folder/f'{task}.json');p(task.capitalize()+':')
        for m in spec['members']:
            c=m['choice'];cfg=configs[c['id']]
            terms=[f"ID {c['id']}",cfg['family'],f"features {cfg['features']}",f"output tau {c['tau']} s",f"parameters {json.dumps(cfg['params'],sort_keys=True)}"]
            for k in ['offset','correction','threshold','threshold_policy']:
                if k in c:terms.append(f'{k} {c[k]}')
            p('; '.join(terms)+'.')
p('The final causal and current-frame manifests retain the complete list of training recordings, preprocessing floors, lag, model filenames, and checksums. For an exact rerun, use those records together with the frozen source, library versions, and original immutable feature store.')
page('Appendix B Reproduction and evidence files')
p('Run the following from calibration_gui using the recorded environment and a new output directory. Reproduction consumes the saved manual feature store and baseline lag decisions. It should not overwrite the original evidence folders.')
for c in commands:p(c)
p('The report generation command produces structured report evidence. The old interactive HTML export helper depended on a previous plugin runtime and is not required for model reproduction. The existing self-contained HTML report remains a readable historical artifact.')
h('Reviewing the evidence without retraining',2)
p('Start with the two model_comparison.csv files for the fixed and tuned outer results. Use before_after.csv to compare the selected procedures. Match fold_metrics.csv by task, model, and outer_fold; match the baseline roi_metrics.csv by task, model, and ROI because it contains several models for each region. The tuned ROI file contains the selected procedure only. Use the force-bin files for conditional errors and paired_uncertainty.json for descriptive group changes.')
p('Read protocol.json before interpreting a candidate ID. Final JSON files under models and models_frame_only bind the selected members to saved estimator files and their checksums. Source manifests identify the exact inputs to the report. The outer prediction parquet files in the original run folders support a full independent metric recalculation, while the smaller copied evidence folder is intended for source review.')
h('Abbreviations',2)
p('ROI means region of interest; MAE, mean absolute error; RMSE, root mean squared error; FPR, false-positive rate; NNLS, nonnegative least squares; SVC, support-vector classification; SVR, support-vector regression; HistGB, histogram gradient boosting; MLP, multilayer perceptron; EMA, exponential moving average; OOF, out-of-fold; and pp, percentage points. Extra Trees refers to extremely randomized trees. XGBoost is the gradient-boosted tree implementation used in the saved environment.')

DOCX=OUT/'Manual_Calibration_Models_Technical_Report.docx'
doc.save(DOCX)
(OUT/'Manual_Calibration_Models_Technical_Report.txt').write_text('\n'.join(text),encoding='utf-8')
prompt='''I am writing a thesis about an optical sensing skin with force estimation, binary contact detection, and nine-ROI localization. Use the attached Manual Calibration Models Technical Report as the primary factual source for the model experiments.\n\nWrite thesis prose in a formal, clear academic style. Begin with an outline, then develop methodology, results, discussion, limitations, conclusions, and recommendations. Keep the three prediction tasks distinct. Use the standard chapter structure unless I supply my university template.\n\nPreserve the manual-only data policy, the 83-recording archive composition, the 54 primary and 11 no-contact eligible recordings, the excluded replay recordings, the exploratory 0.05–3 N interval, the six whole-TEST-group outer folds, train-only preprocessing and inner tuning, and equal-recording metric weights. Distinguish the earlier fixed-family winners, earlier nested selector, tuned nested selector, retrospectively ranked tuned families, and final all-data refits. Do not report final training-selection scores as independent test results.\n\nUse the report's exact metrics and units. Explain percentage-point versus relative changes. Main tuned selector results are force MAE 0.499 N and RMSE 0.632 N, localization accuracy 84.40% and macro F1 84.75%, and contact balanced accuracy 98.50%, recall 97.00%, and zero observed FPR. Describe these as retrospective results on this dataset. Do not imply uniform improvement: force MAE worsened in all four bins below 2 N, ROI 2 recall declined, and TEST1 remains weak.\n\nDo not claim prospective validation, validated detection limit or force range, continuous spatial error in millimetres, calibrated contact probabilities, multi-contact force accuracy, zero real-world false alarms, deployed live latency, or integration of these models into the GUI. Causal history can delay response; stationary-ROI archived sessions do not establish dynamic localization performance.\n\nUse project source IDs for traceability during drafting, and flag where formal external literature citations are needed. Do not invent papers, DOI values, apparatus dimensions, load-cell specifications, experimental conditions, uncertainty budgets, or missing laboratory procedures. Ask me for missing lab facts, or use clearly marked placeholders. Treat historical hardware tests separately from the model evaluation. Explain that nested validation does not undo previous retrospective use of the data.\n\nProduce interpretable tables with metric definitions, sample counts, comparison procedures, and units. Retain the limitations in the abstract, discussion, and conclusions at an appropriate level. Do not claim a single technique caused the gains without a controlled ablation. Conclude with a prospective manual-only validation plan and the outstanding force-model limitations.\n'''
(OUT/'ChatGPT_Writing_Prompt.txt').write_text(prompt,encoding='utf-8')
evidence=OUT/'evidence';evidence.mkdir(exist_ok=True)
manifest=[]
files=['protocol.json','model_comparison.csv','fold_metrics.csv','roi_metrics.csv','force_bin_metrics.csv','end_to_end_metrics.json','validation_audit.json','provenance.json']
for run in [B,T]:
    for name in files+(['evaluation_coverage.csv','estimator_parameters.json','lag_selection.json'] if run==B else ['before_after.csv','paired_uncertainty.json','frame_only_protocol.json','inference_timing.json','environment.json','final_selections.json','report_delivery.json','source_snapshot_manifest.json']):
        src=run/name
        if not src.exists():continue
        dest=evidence/run.name/name;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dest)
        manifest.append({'source':str(src.relative_to(ROOT)),'copy':str(dest.relative_to(OUT)),'sha256':hashlib.sha256(src.read_bytes()).hexdigest()})
for folder in ['models','models_frame_only']:
    for src in (T/folder).glob('*.json'):
        dest=evidence/T.name/folder/src.name;dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dest)
        manifest.append({'source':str(src.relative_to(ROOT)),'copy':str(dest.relative_to(OUT)),'sha256':hashlib.sha256(src.read_bytes()).hexdigest()})
(evidence/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
assert len(protocol['candidates'])==178
assert Counter(c['task'] for c in protocol['candidates'])=={'force':62,'localization':64,'contact':52}
assert sum(int(r['evaluated_frames']) for r in coverage)==7154
assert audit['status']=='passed' and audit['metric_checks']==114
(OUT/'report_build_checks.json').write_text(json.dumps({'source_checks':'passed','tables':tables,'text_words':len(' '.join(text).split()),'evidence_files':len(manifest),'source_audit_status':audit['status'],'visual_review':'pending'},indent=2),encoding='utf-8')
print(json.dumps({'docx':str(DOCX),'tables':tables,'words':len(' '.join(text).split()),'evidence_files':len(manifest)}))
